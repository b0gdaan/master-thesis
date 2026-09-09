from pathlib import Path
import argparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT / "govor-na-srpskom.md"
DEFAULT_OUTPUT = ROOT / "Tekst_izlaganja_Bogdan_Babaev.docx"


def set_run_font(run, size=None, bold=None, italic=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:cs"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_markdown_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.15

    bold = False
    parts = text.split("**")
    for index, part in enumerate(parts):
        if not part:
            bold = not bold
            continue
        run = paragraph.add_run(part)
        set_run_font(run, 12, bold=bold)
        bold = not bold
    return paragraph


def normalize_math(text):
    replacements = {
        "$R^2$": "R²",
        "$-0,75 \\sigma_t$": "-0,75 σₜ",
        "$(-1,1)$": "(-1, 1)",
        "$w = 14, 30, 60, 90$": "w = 14, 30, 60, 90",
        "$w=30$": "w = 30",
        "$w = 30$": "w = 30",
        "$\\times$": "x",
        "DCC-GARCH(1,1)": "DCC-GARCH(1,1)",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def configure_style(style, size, bold=False):
    style.font.name = "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    style._element.rPr.rFonts.set(qn("w:cs"), "Aptos")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--title", default="Текст излагања за одбрану мастер рада")
    args = parser.parse_args()

    source = args.source if args.source.is_absolute() else ROOT / args.source
    output = args.output if args.output.is_absolute() else ROOT / args.output

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    configure_style(doc.styles["Normal"], 12)
    configure_style(doc.styles["Title"], 19, bold=True)
    configure_style(doc.styles["Heading 1"], 14, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_run_font(title.add_run(args.title), 19, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run(
            "Предвиђање временски променљивих међутржишних зависности "
            "између криптовалута и традиционалне имовине применом машинског учења\n"
            "Богдан Бабаев"
        ),
        11,
        italic=True,
    )

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = normalize_math(raw_line.strip())
        if not line or line.startswith("# Текст") or line.startswith("Текст је написан"):
            continue
        if line.startswith("## "):
            heading = doc.add_paragraph(style="Heading 1")
            heading.paragraph_format.space_before = Pt(14)
            heading.paragraph_format.space_after = Pt(7)
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.keep_together = True
            set_run_font(heading.add_run(line[3:]), 14, bold=True)
        else:
            add_markdown_paragraph(doc, line)

    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
